"""Tạo prompt.docx (mô tả yêu cầu phần mềm). Chạy: python tools/make_prompt_docx.py"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "prompt.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(11)

title = doc.add_heading("PROMPT: Xây phần mềm web phân tích chứng khoán VN50 & Ngân hàng", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta = doc.add_paragraph("Phiên bản cập nhật 07/10/2026 · Repo: https://github.com/tieunguyetpham/Claude-CK · "
                         "Triển khai: https://trading.tieunguyetpham.store")
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.runs[0].font.size = Pt(9)
meta.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def h(text):
    doc.add_heading(text, level=1)


def p(text):
    doc.add_paragraph(text)


def bullets(items, style="List Bullet"):
    for it in items:
        doc.add_paragraph(it, style=style)


def code(text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)


p("Hãy viết cho tôi một ứng dụng web dashboard bằng Python + Streamlit để phân tích cổ phiếu Việt Nam. "
  "Code sạch, chia module rõ ràng, chú thích tiếng Việt ở những chỗ quan trọng.")
para = doc.add_paragraph()
para.add_run("Ba ưu tiên xuyên suốt: ").bold = True
para.add_run("(1) giao diện đơn giản, thiết kế gọn; (2) ổn định, giảm thiểu lỗi; "
             "(3) phân tích thị trường kịp thời — dữ liệu mới, hiển thị nhanh.")

h("0. Nguyên tắc làm việc (bắt buộc)")
bullets([
    "Làm từng bước. Gặp lỗi ở bước nào phải xử lý dứt điểm toàn bộ lỗi của bước đó rồi mới sang bước tiếp theo.",
    "Kiểm tra đầu vào trước khi viết phần mềm: xác minh API dữ liệu thật sự trả dữ liệu (đúng định dạng, đúng đơn vị) "
    "cho toàn bộ danh sách mã, và kiểm thử các kiểu nhập liệu sai của người dùng.",
    "Những gì cần người dùng cài đặt / cung cấp (Python, quyền truy cập VPS, DNS...) phải hỏi gộp một lần ngay từ đầu.",
    "Không giả định thư viện còn hoạt động: thử cài và gọi thật trước khi dùng.",
])

h("1. Mục tiêu")
p("Người dùng chọn hoặc nhập một/nhiều mã cổ phiếu (ví dụ VCB, FPT, HPG, TCB, MWG). Ứng dụng tự động:")
bullets([
    "Lấy dữ liệu giá từ API của sàn/công ty chứng khoán.",
    "Tính các chỉ báo phân tích kỹ thuật.",
    "Vẽ biểu đồ và bảng số liệu.",
    "Đưa ra nhận định tổng hợp MUA / GIỮ / BÁN (mang tính tham khảo) kèm lý do cụ thể.",
], style="List Number")

h("2. Nguồn dữ liệu — gọi trực tiếp API của sàn")
p("Không dùng thư viện vnstock: gói vnstock và phụ thuộc vnai đã bị gỡ khỏi PyPI (vnstock3 không cài được). "
  "Endpoint TCBS cũ (apipubaws.tcbs.com.vn) đã ngừng. Dùng API công khai của SSI iBoard bằng thư viện requests:")
bullets([
    "Lịch sử nến ngày (OHLCV): GET https://iboard-api.ssi.com.vn/statistics/charts/history"
    "?resolution=1D&symbol=VCB&from=<unix>&to=<unix> → data gồm các mảng t, o, h, l, c, v. "
    "Giá tính theo NGHÌN ĐỒNG (57.2 = 57.200 đ).",
    "Báo giá realtime + thông tin mã: GET https://iboard-query.ssi.com.vn/stock/VCB → matchedPrice, priceChange, "
    "priceChangePercent, refPrice, highest, lowest, nmTotalTradedQty, listedShare, companyNameVi, exchange, tradingDate. "
    "Giá tính theo ĐỒNG → chia 1000 để thống nhất đơn vị nghìn đồng.",
    "Gửi header User-Agent và Accept: application/json; timeout ~15 giây; tự thử lại 2 lần khi lỗi mạng.",
    "Mã không tồn tại vẫn trả HTTP 200 nhưng mảng rỗng → phải kiểm tra và báo \"Không tìm thấy dữ liệu cho mã X\".",
    "Ngoài giờ giao dịch matchedPrice có thể trống → dùng refPrice/priorClosePrice.",
    "listedShare = 0 với một số mã → vốn hóa hiển thị \"—\". API công khai không có P/E, P/B, EPS, ROE → không hiển thị.",
    "Lấy khoảng 3 năm lịch sử để MA200 đủ dữ liệu, chỉ cắt theo khoảng thời gian người dùng chọn khi vẽ.",
    "Ghép báo giá realtime vào nến cuối (thêm nến phiên hôm nay hoặc cập nhật nến cùng ngày) để phân tích luôn dùng "
    "giá mới nhất trong phiên, khớp với giá đang hiển thị.",
    "Cache bằng st.cache_data (TTL 5 phút, giới hạn 64 mục để không phình RAM khi nhiều người dùng) nhưng KHÔNG lưu "
    "cache kết quả lỗi mạng (lỗi tạm thời phải thử lại ngay).",
    "Tải nhiều mã song song bằng ThreadPoolExecutor; mỗi luồng dùng requests.Session riêng (threading.local).",
])

h("3. Danh sách mã")
bullets([
    "Rổ VN50 = chỉ số VNX50 (50 mã lớn nhất HOSE + HNX). SSI không có nhóm tên \"VN50\" (trả rỗng). Lấy thành phần "
    "trực tiếp: GET https://iboard-query.ssi.com.vn/stock/group/VNX50 → data[].stockSymbol; cache 6 giờ để tự cập nhật "
    "khi rổ cơ cấu lại. Kèm danh sách dự phòng (snapshot) khi API lỗi; không lưu cache khi phải dùng dự phòng.",
    "Lọc kết quả rổ: chỉ nhận chuỗi 3–10 ký tự A–Z/0–9 (giá trị None không được thành mã \"NONE\"); dưới 20 mã coi như lỗi.",
    "Không tự soạn tay danh sách rổ chỉ số: danh sách tự soạn dễ lỗi thời (đã kiểm chứng lệch 11/50 mã so với VNX50).",
    "Nhóm ngân hàng (27 mã, HOSE/HNX/UPCoM): VCB, BID, CTG, TCB, MBB, ACB, VPB, HDB, STB, TPB, SHB, VIB, LPB, SSB, "
    "EIB, MSB, OCB, NAB, ABB, BVB, KLB, VAB, BAB, SGB, PGB, NVB, VBB.",
    "Mã mặc định trong ô chọn phải thuộc danh sách lựa chọn (nếu không Streamlit báo lỗi).",
    "Trước khi dùng, chạy kiểm tra để chắc chắn mọi mã trong danh sách đều lấy được dữ liệu mới nhất.",
])

h("4. Kiểm tra đầu vào người dùng")
bullets([
    "Gộp mã chọn từ danh sách + mã gõ tay; tách bằng dấu phẩy, chấm phẩy hoặc khoảng trắng.",
    "Chuẩn hóa chữ hoa, bỏ trùng; chỉ chấp nhận 3–10 ký tự A–Z, 0–9; mã không hợp lệ báo rõ và bỏ qua.",
    "Tối đa 10 mã mỗi lần để giữ tốc độ; không chọn mã nào thì hiện hướng dẫn thay vì lỗi.",
    "Dữ liệu bất thường (mã mới niêm yết ít phiên, giá đứng yên, khối lượng 0) không được làm sập ứng dụng.",
])

h("5. Giao diện (Streamlit)")
bullets([
    "Thanh bên: multiselect rổ VN50 & ngân hàng (đánh dấu mã ngân hàng), ô nhập mã khác, chọn khoảng "
    "3 tháng / 6 tháng / 1 năm / 2 năm, nút \"Làm mới dữ liệu\".",
    "Đầu trang: tiêu đề, cảnh báo miễn trừ trách nhiệm, thời điểm lấy dữ liệu + nguồn.",
    "Nhiều mã: bảng so sánh nhanh, cột quan trọng trước để không bị khuất trên điện thoại: mã, nhận định, giá, "
    "% thay đổi, điểm, RSI.",
    "Ô chỉ số không được tràn chữ ở mọi kích thước (điện thoại 375px, máy tính bảng 800px, máy tính 1400px): khối lượng "
    "dạng gọn \"13.40 tr\" (số đầy đủ trong chú thích), giá thấp phiên ở dòng phụ, cỡ chữ co theo độ rộng ô (CSS "
    "container query). Biểu đồ: chú thích không đè thanh công cụ, ngày dd/mm/yyyy, số làm tròn khi rê chuột.",
    "Mỗi mã một tab: thẻ tóm tắt (giá, %, khối lượng, cao/thấp, vốn hóa) → hộp nhận định (xanh MUA / vàng GIỮ / đỏ BÁN) "
    "kèm lý do → biểu đồ → bảng chỉ báo, hỗ trợ/kháng cự, nút tải CSV (utf-8-sig để Excel đọc đúng tiếng Việt).",
    "Biểu đồ Plotly 4 tầng: nến + MA20/50/200 + Bollinger + đường hỗ trợ/kháng cự; khối lượng; RSI (mốc 30/70); "
    "MACD. Ẩn ngày nghỉ (rangebreaks) để nến liền mạch. Màu tăng xanh, giảm đỏ.",
    "Mọi chữ trên giao diện bằng tiếng Việt.",
])

h("6. Chỉ báo kỹ thuật (tự tính bằng pandas/numpy, không dùng pandas-ta)")
bullets([
    "SMA 20/50/200, EMA20.",
    "RSI(14) kiểu Wilder — giá đứng yên trả 50, chỉ tăng trả 100.",
    "MACD(12, 26, 9), Bollinger Bands(20, 2), Stochastic(14, 3), ADX(14) kèm +DI/−DI.",
    "Khối lượng trung bình 20 phiên; hỗ trợ/kháng cự = đáy/đỉnh 20 và 60 phiên gần nhất.",
])

h("7. Logic nhận định")
bullets([
    "Mỗi tín hiệu cho +1 (tích cực), −1 (tiêu cực), 0 (trung tính): giá so với MA50; MA20 so với MA50 (ưu tiên giao cắt "
    "trong 3 phiên); giá so với MA200; RSI (<30 quá bán, >70 quá mua); MACD so với signal (ưu tiên giao cắt); "
    "Bollinger (thủng dải dưới / vượt dải trên); Stochastic (vùng quá bán/quá mua có giao cắt); khối lượng đột biến "
    ">1,5 lần TB20 theo chiều giá.",
    "Tổng điểm ≥ +3 → MUA; ≤ −3 → BÁN; còn lại → GIỮ. Hiển thị lý do từng tín hiệu.",
    "ADX ≥ 25: xu hướng mạnh, tín hiệu đáng tin hơn; < 25: xu hướng yếu/đi ngang, nên thận trọng (không cộng điểm).",
    "Cổ phiếu ngân hàng: ghi chú định giá theo P/B, ROE; nhạy với lãi suất, tăng trưởng tín dụng, nợ xấu.",
])

h("8. Ràng buộc")
bullets([
    "Không đặt lệnh mua/bán, không kết nối tài khoản giao dịch thật — chỉ là công cụ phân tích tham khảo.",
    "Luôn hiển thị: \"Thông tin chỉ mang tính tham khảo, không phải khuyến nghị đầu tư. Nhà đầu tư tự chịu trách "
    "nhiệm với quyết định của mình.\"",
    "Mọi lời gọi API/tính toán bọc try/except; một mã lỗi không được làm hỏng các mã khác.",
])

h("9. Cấu trúc dự án")
code("app.py            giao diện Streamlit\n"
     "data.py           gọi API SSI, thử lại, chuẩn hóa đơn vị\n"
     "indicators.py     chỉ báo kỹ thuật\n"
     "analysis.py       chấm điểm & nhận định\n"
     "symbols.py        danh sách VN50/ngân hàng, kiểm tra mã nhập\n"
     "tests/            pytest cho đầu vào và dữ liệu bất thường\n"
     "requirements.txt  khóa phiên bản đã kiểm thử (Python >= 3.11)\n"
     "Dockerfile        đóng gói app (python:3.12-slim, user không phải root, healthcheck)\n"
     "deploy/           docker-compose.yml, nginx-trading.conf, deploy.sh\n"
     ".gitignore, .dockerignore, .gitattributes (ép LF cho .sh/.conf), README.md")

h("10. Kiểm thử trước khi bàn giao")
bullets([
    "Quét toàn bộ mã VNX50 + ngân hàng: lấy được dữ liệu, ngày mới nhất, giá realtime khớp biểu đồ, không có giá ≤ 0.",
    "Giả lập tình huống biên trên giao diện: báo giá lỗi nhưng có lịch sử; mã mới niêm yết 1, 2, 30 phiên; không lấy "
    "được thành phần rổ (dùng dự phòng, lần sau thử lại).",
    "pytest cho: chuỗi rỗng, chữ thường, trùng mã, ký tự đặc biệt, chữ có dấu, chuỗi dài; dữ liệu 1 phiên, 2 phiên, "
    "giá đứng yên, chỉ tăng, chỉ giảm, khối lượng 0; ghép báo giá realtime (phiên mới, cùng phiên, báo giá cũ, "
    "chưa khớp lệnh, ngày sai định dạng).",
    "Giả lập mất mạng: app báo lỗi rõ ràng; có mạng lại thì tải được ngay, không bị kẹt lỗi trong cache.",
    "streamlit.testing AppTest cho các kịch bản giao diện (mặc định, mã sai, không chọn mã, quá 10 mã, đổi khoảng thời gian).",
    "Chạy app thật, kiểm tra trình duyệt không có lỗi console, log server sạch.",
])

h("11. Lưu code trên GitHub")
bullets([
    "Đẩy code lên https://github.com/tieunguyetpham/Claude-CK (nhánh main).",
    "Không commit venv, file .env, khóa bí mật, file CSV xuất ra.",
])

h("12. Triển khai trên VPS")
p("App phục vụ tại tên miền con https://trading.tieunguyetpham.store trên VPS 116.118.6.222 (Ubuntu 24.04). "
  "VPS đang chạy hệ thống ppmeeting bằng Docker — tuyệt đối không làm gián đoạn hệ thống này.")
bullets([
    "Khảo sát VPS trước (chỉ đọc): hệ điều hành, RAM, cổng đang dùng, web server đang giữ cổng 80/443, "
    "cấu hình Nginx, chứng chỉ và cơ chế gia hạn hiện có. Không cài thêm web server tranh cổng 80/443.",
    "Hiện trạng: container ppmeeting-nginx-1 (nginx:1.27-alpine) giữ cổng 80/443, cấu hình tại "
    "/opt/ppmeeting/infra/nginx/nginx.conf, chứng chỉ tại /opt/ppmeeting/infra/nginx/certs; certbot cấp kiểu "
    "standalone, có hook pre/post (dừng/chạy lại Nginx) và deploy (chép chứng chỉ) tại /etc/letsencrypt/renewal-hooks.",
    "Đóng gói app bằng Dockerfile; chạy container riêng trading-app qua deploy/docker-compose.yml, giới hạn 512 MB RAM, "
    "restart unless-stopped, healthcheck /_stcore/health, nối vào mạng có sẵn ppmeeting_default (external).",
    "DNS: bản ghi A trading.tieunguyetpham.store → 116.118.6.222.",
    "Repo GitHub để Public (đã rà soát không có bí mật, email tác giả dùng địa chỉ ẩn danh noreply của GitHub); "
    "VPS git clone qua HTTPS vào /opt/Claude-CK, không cần khóa truy cập.",
    "HTTPS: bổ sung trading.tieunguyetpham.store vào chứng chỉ sẵn có (certbot certonly --standalone --expand "
    "--cert-name tieunguyetpham.store) để cơ chế tự gia hạn tiếp tục hoạt động. Chạy hook pre/deploy/post của ppmeeting "
    "thủ công, không truyền --pre-hook/--post-hook (tránh certbot lưu hook trùng); hook post luôn chạy kể cả khi lỗi.",
    "Nginx: sao lưu nginx.conf, thêm một server block server_name trading.tieunguyetpham.store dùng chung chứng chỉ; "
    "proxy tới http://trading-app:8501 qua biến và resolver 127.0.0.11 (Nginx vẫn khởi động khi trading-app dừng); "
    "hỗ trợ WebSocket. nginx.conf được bind-mount dạng một file đơn → sửa bằng cách ghi đè nội dung (giữ inode), "
    "không dùng sed -i/mv. Thử nginx -t trên container tạm trước, rồi mới nginx -t và nginx -s reload thật.",
    "deploy/deploy.sh dùng để cập nhật: git pull → docker compose up -d --build → chờ healthy → chỉ dọn image cũ "
    "có nhãn app=trading-app (không đụng image của ppmeeting).",
    "Kiểm tra sau triển khai: trang trading chạy qua HTTPS, trang ppmeeting và /api/ vẫn hoạt động bình thường.",
])

doc.save(OUT)
print("saved", OUT)
